#!/usr/bin/env python3
"""Audit current location coverage or build a registry from reviewed evidence."""

import argparse
import csv
import json
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

BACKEND_ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = BACKEND_ROOT / "data" / "Copy of crop_recommendation_45_with_year_area.csv"
CURRENT_REGISTRY_PATH = BACKEND_ROOT / "data" / "weather_location_registry.json"
AUDIT_PATH = BACKEND_ROOT / "data" / "weather_location_resolution_audit.json"
INPUT_TEMPLATE_DIR = BACKEND_ROOT / "data" / "weather_location_resolution"
LGD_REPORT_INPUT_DIR = INPUT_TEMPLATE_DIR / "lgd_reports"
BATCH_AUDIT_PATH = INPUT_TEMPLATE_DIR / "lgd_batch_resolution_audit.json"
LGD_DOWNLOAD_URL = "https://lgdirectory.gov.in/downloadDirectory.do"

LGD_FIELDS = {
    "state_lgd_code",
    "district_lgd_code",
    "district_version",
    "source_record_id",
    "state_name",
    "district_name",
    "source_url",
    "source_downloaded_at",
    "source_report_generated_at",
    "source_file",
    "source_worksheet",
    "source_worksheet_row",
}
CROSSWALK_FIELDS = {
    "crop_state",
    "crop_district",
    "state_lgd_code",
    "district_lgd_code",
    "official_lgd_district_name",
    "district_version",
    "source_record_id",
    "match_method",
    "review_status",
    "evidence_url",
    "review_note",
    "reviewed_at",
    "reviewed_by",
}
COORDINATE_FIELDS = {
    "state_lgd_code",
    "district_lgd_code",
    "latitude",
    "longitude",
    "reference_place",
    "reference_type",
    "coordinate_source",
    "coordinate_source_url",
    "coordinate_source_record",
    "verification_source_url",
    "verification_status",
    "verification_note",
    "verified_at",
    "verified_by",
}


def read_csv(path: Path, required_fields: set[str]) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        missing = required_fields - set(reader.fieldnames or ())
        if missing:
            raise ValueError(f"{path} is missing required columns: {sorted(missing)}")
        return [{key: (value or "").strip() for key, value in row.items()} for row in reader]


def name_key(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip()).casefold()


def _excel_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    return str(value).strip()


def _excel_column_name(column_index: int) -> str:
    column_name = ""
    while column_index:
        column_index, remainder = divmod(column_index - 1, 26)
        column_name = chr(65 + remainder) + column_name
    return column_name


def _load_spreadsheet_rows(path: Path) -> list[tuple[str, list[list[Any]]]]:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "Reading .xlsx reports requires the openpyxl dependency."
            ) from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                (
                    worksheet.title,
                    [list(row) for row in worksheet.iter_rows(values_only=True)],
                )
                for worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError(
                "Reading .xls reports requires the xlrd dependency."
            ) from exc
        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            return [
                (
                    worksheet.name,
                    [
                        list(worksheet.row_values(row_index))
                        for row_index in range(worksheet.nrows)
                    ],
                )
                for worksheet in workbook.sheets()
            ]
        finally:
            workbook.release_resources()
    raise ValueError(f"Unsupported LGD report extension: {path.suffix}")


def _parse_lgd_report(
    path: Path, downloaded_at: str
) -> tuple[str, str, list[dict[str, str]]]:
    """Parse one LGD workbook only when its report title identifies its state."""
    title_pattern = re.compile(
        r"All\s+Districts\s+of\s+(.+?)\s*\(\s*State\s*Code\s*:\s*(\d+)\s*\)\s*State?",
        re.IGNORECASE,
    )
    workbook_sheets = _load_spreadsheet_rows(path)
    report_identities: set[tuple[str, str]] = set()
    generated_at: str | None = None
    generated_formats = (
        "%b %d, %Y, %I:%M %p",
        "%B %d, %Y, %I:%M %p",
        "%d-%m-%Y %H:%M",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
    )
    for _, rows in workbook_sheets:
        for row in rows:
            for value in row:
                text = _excel_cell_text(value)
                match = title_pattern.search(text)
                if match:
                    report_identities.add(
                        (re.sub(r"\s+", " ", match.group(1)).strip(), match.group(2))
                    )
                if isinstance(value, datetime):
                    candidate_time = value.isoformat(sep=" ", timespec="minutes")
                else:
                    candidate_time = text
                for time_format in generated_formats:
                    try:
                        generated_at = datetime.strptime(
                            candidate_time, time_format
                        ).isoformat(sep=" ", timespec="minutes")
                        break
                    except ValueError:
                        continue
    if len(report_identities) != 1:
        raise ValueError(
            "Report must contain exactly one identifiable 'All Districts of "
            "<state> (State Code:<code>) State' title."
        )
    if generated_at is None:
        raise ValueError(
            "Report does not contain a recognizable generated timestamp."
        )
    state_name, state_code = next(iter(report_identities))

    parsed_records: list[dict[str, str]] = []
    for sheet_name, rows in workbook_sheets:
        for header_index, row in enumerate(rows):
            normalized = [name_key(_excel_cell_text(value)) for value in row]
            code_columns = [
                index for index, value in enumerate(normalized)
                if value == "district code"
            ]
            version_columns = [
                index for index, value in enumerate(normalized)
                if value == "district version"
            ]
            if not code_columns or not version_columns:
                continue
            district_name_column: int | None = None
            header_end = header_index
            for index, value in enumerate(normalized):
                combined = value
                if header_index + 1 < len(rows) and index < len(rows[header_index + 1]):
                    combined += " " + name_key(
                        _excel_cell_text(rows[header_index + 1][index])
                    )
                if "district name" in combined and (
                    "english" in combined
                    or not any("district name" in candidate for candidate in normalized)
                ):
                    district_name_column = index
                    if "english" in combined:
                        header_end = header_index + 1
                    break
            if district_name_column is None:
                district_name_candidates = [
                    index for index, value in enumerate(normalized)
                    if "district name" in value
                ]
                if district_name_candidates:
                    district_name_column = district_name_candidates[0]
            if district_name_column is None:
                continue

            code_column = code_columns[0]
            version_column = version_columns[0]
            for row_index in range(header_end + 1, len(rows)):
                row_values = rows[row_index]
                required_index = max(
                    code_column, version_column, district_name_column
                )
                if len(row_values) <= required_index:
                    continue
                district_code = _excel_cell_text(row_values[code_column])
                district_version = _excel_cell_text(row_values[version_column])
                district_name = _excel_cell_text(row_values[district_name_column])
                if not (district_code and district_version and district_name):
                    continue
                if district_code.endswith(".0") and district_code[:-2].isdigit():
                    district_code = district_code[:-2]
                if district_version.endswith(".0") and district_version[:-2].isdigit():
                    district_version = district_version[:-2]
                excel_row = row_index + 1
                parsed_records.append(
                    {
                        "state_lgd_code": state_code,
                        "district_lgd_code": district_code,
                        "district_version": district_version,
                        "source_record_id": (
                            f"{path.name}:{sheet_name}!"
                            f"{_excel_column_name(code_column + 1)}{excel_row}:"
                            f"{_excel_column_name(district_name_column + 1)}{excel_row}"
                        ),
                        "state_name": state_name,
                        "district_name": district_name,
                        "source_url": LGD_DOWNLOAD_URL,
                        "source_downloaded_at": downloaded_at,
                        "source_report_generated_at": generated_at,
                        "source_file": path.name,
                        "source_worksheet": sheet_name,
                        "source_worksheet_row": str(excel_row),
                    }
                )
    if not parsed_records:
        raise ValueError("No district rows were found under a recognized LGD header.")
    return state_name, state_code, parsed_records


def _write_lgd_snapshot(path: Path, records: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=sorted(LGD_FIELDS))
        writer.writeheader()
        writer.writerows(records)


def _state_snapshot_audit(
    state: str,
    pairs: list[tuple[str, str]],
    records: list[dict[str, str]],
    crosswalk_by_pair: dict[tuple[str, str], list[dict[str, str]]],
    report_status: str,
    conflicts: list[str],
) -> dict[str, Any]:
    names = {name_key(row["district_name"]): row for row in records}
    direct_matches = [
        {
            "crop_district": district,
            "official_lgd_district_name": names[name_key(district)][
                "district_name"
            ],
            "district_lgd_code": names[name_key(district)]["district_lgd_code"],
            "district_version": names[name_key(district)]["district_version"],
        }
        for pair_state, district in pairs
        if pair_state == state and name_key(district) in names
    ]
    reviewed_aliases = []
    lgd_rows_by_id = {
        (row["state_lgd_code"], row["district_lgd_code"]): row
        for row in records
    }
    for pair_state, district in pairs:
        if pair_state != state or name_key(district) in names:
            continue
        mappings = crosswalk_by_pair.get((state, district), [])
        mapping = mappings[0] if len(mappings) == 1 else None
        lgd_row = lgd_rows_by_id.get(
            (mapping["state_lgd_code"], mapping["district_lgd_code"])
        ) if mapping else None
        if (
            mapping
            and lgd_row
            and mapping["review_status"].casefold() == "approved"
            and mapping["match_method"].casefold()
            in {"reviewed_alias", "historical_rename"}
            and name_key(mapping["official_lgd_district_name"])
            == name_key(lgd_row["district_name"])
            and mapping["district_version"] == lgd_row["district_version"]
            and mapping["source_record_id"] == lgd_row["source_record_id"]
            and mapping["evidence_url"]
            and mapping["review_note"]
        ):
            reviewed_aliases.append(
                {
                    "crop_district": district,
                    "official_lgd_district_name": mapping[
                        "official_lgd_district_name"
                    ],
                    "district_lgd_code": mapping["district_lgd_code"],
                    "reviewed_at": mapping["reviewed_at"],
                    "reviewed_by": mapping["reviewed_by"],
                    "evidence_url": mapping["evidence_url"],
                }
            )
    matched_labels = {
        name_key(row["crop_district"]) for row in direct_matches
    } | {
        name_key(row["crop_district"]) for row in reviewed_aliases
    }
    unresolved = [
        district
        for pair_state, district in pairs
        if pair_state == state and name_key(district) not in matched_labels
    ]
    return {
        "report_status": report_status,
        "crop_pair_count": sum(pair_state == state for pair_state, _ in pairs),
        "lgd_record_count": len(records),
        "direct_matches": direct_matches,
        "reviewed_aliases": reviewed_aliases,
        "unresolved_names": unresolved,
        "conflicting_records": conflicts,
    }


def audit_lgd_reports(
    report_dir: Path = LGD_REPORT_INPUT_DIR,
    output_dir: Path = INPUT_TEMPLATE_DIR,
    csv_path: Path = CSV_PATH,
    crosswalk_path: Path = INPUT_TEMPLATE_DIR / "crop_to_lgd_crosswalk.csv",
) -> dict[str, Any]:
    """Import identifiable LGD workbooks and audit every crop state label."""
    pairs, _, source_errors = crop_pairs(csv_path)
    if source_errors:
        raise ValueError("Crop CSV contains missing location labels: " + "; ".join(source_errors[:5]))
    states = sorted({state for state, _ in pairs})
    crosswalk_rows = read_csv(crosswalk_path, CROSSWALK_FIELDS)
    crosswalk_by_pair: dict[
        tuple[str, str], list[dict[str, str]]
    ] = defaultdict(list)
    for row in crosswalk_rows:
        crosswalk_by_pair[(row["crop_state"], row["crop_district"])].append(row)
    report_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    candidates: dict[str, list[dict[str, Any]]] = defaultdict(list)
    identity_conflicts: list[dict[str, str]] = []
    for snapshot_path in sorted(output_dir.glob("lgd_district_snapshot_*.csv")):
        if snapshot_path.name == "lgd_district_snapshot_template.csv":
            continue
        snapshot_rows = read_csv(snapshot_path, LGD_FIELDS)
        snapshot_states = {row["state_name"] for row in snapshot_rows}
        if len(snapshot_states) != 1:
            identity_conflicts.append(
                {
                    "file": snapshot_path.name,
                    "reason": "Snapshot contains records for multiple/no state names.",
                }
            )
            continue
        snapshot_state = next(iter(snapshot_states))
        candidates[snapshot_state].append(
            {"path": snapshot_path, "records": snapshot_rows, "source": "snapshot"}
        )

    rejected_reports: list[dict[str, str]] = []
    state_by_normalized_name = {name_key(state): state for state in states}
    for report_path in sorted(report_dir.iterdir()):
        if report_path.name.startswith("~$") or not report_path.is_file():
            continue
        if report_path.suffix.casefold() not in {".xls", ".xlsx"}:
            continue
        try:
            downloaded_at = datetime.fromtimestamp(
                report_path.stat().st_mtime
            ).date().isoformat()
            report_state, _, records = _parse_lgd_report(
                report_path, downloaded_at
            )
        except (OSError, ValueError, RuntimeError) as exc:
            rejected_reports.append(
                {"file": report_path.name, "reason": str(exc)}
            )
            continue
        canonical_state = state_by_normalized_name.get(name_key(report_state))
        if canonical_state is None:
            rejected_reports.append(
                {
                    "file": report_path.name,
                    "reason": (
                        f"Report identifies state {report_state!r}, which is not "
                        "an exact crop CSV state label; file was not imported."
                    ),
                }
            )
            continue
        for record in records:
            record["state_name"] = canonical_state
        candidates[canonical_state].append(
            {"path": report_path, "records": records, "source": "workbook"}
        )
        duplicates_of_same_source = [
            candidate
            for candidate in candidates[canonical_state]
            if candidate["source"] == "snapshot"
            and all(
                row["source_file"] == report_path.name
                for row in candidate["records"]
            )
        ]
        for candidate in duplicates_of_same_source:
            existing_identity = {
                (
                    row["state_lgd_code"],
                    row["district_lgd_code"],
                    row["district_version"],
                    name_key(row["district_name"]),
                )
                for row in candidate["records"]
            }
            workbook_identity = {
                (
                    row["state_lgd_code"],
                    row["district_lgd_code"],
                    row["district_version"],
                    name_key(row["district_name"]),
                )
                for row in records
            }
            if existing_identity == workbook_identity:
                candidates[canonical_state].remove(candidate)

    per_state: dict[str, Any] = {}
    imported_snapshots: list[str] = []
    missing_reports: list[str] = []
    for state in states:
        state_candidates = candidates.get(state, [])
        conflicts: list[str] = []
        if len(state_candidates) > 1:
            conflicts.extend(
                f"Multiple state snapshots/reports claim this state: "
                f"{', '.join(str(item['path'].name) for item in state_candidates)}."
            )
        selected = state_candidates[0] if len(state_candidates) == 1 else None
        records = selected["records"] if selected else []
        if selected:
            if any(row["state_name"] != state for row in records):
                conflicts.append("Snapshot row state name does not match expected state.")
            state_codes = {row["state_lgd_code"] for row in records}
            if len(state_codes) != 1 or not next(iter(state_codes), ""):
                conflicts.append(
                    "Snapshot rows do not contain exactly one non-empty LGD state code."
                )
            duplicate_codes = [
                code for code, count in Counter(
                    row["district_lgd_code"] for row in records
                ).items() if count > 1
            ]
            duplicate_names = [
                name for name, count in Counter(
                    name_key(row["district_name"]) for row in records
                ).items() if count > 1
            ]
            if duplicate_codes:
                conflicts.append(
                    "Duplicate district LGD codes: " + ", ".join(duplicate_codes)
                )
            if duplicate_names:
                conflicts.append(
                    "Duplicate normalized district names: "
                    + ", ".join(duplicate_names)
                )
            if selected["source"] == "workbook" and not conflicts:
                snapshot_name = (
                    "lgd_district_snapshot_"
                    + re.sub(r"[^a-z0-9]+", "_", state.casefold()).strip("_")
                    + ".csv"
                )
                snapshot_path = output_dir / snapshot_name
                _write_lgd_snapshot(snapshot_path, records)
                imported_snapshots.append(str(snapshot_path))
        if not selected:
            missing_reports.append(state)
        status = (
            "conflicting_records"
            if conflicts
            else "missing_report"
            if not selected
            else "processed"
        )
        state_audit = _state_snapshot_audit(
            state,
            pairs,
            records if selected and not conflicts else [],
            crosswalk_by_pair,
            status,
            conflicts,
        )
        duplicate_crosswalk_pairs = [
            district
            for pair_state, district in pairs
            if pair_state == state
            and len(crosswalk_by_pair.get((state, district), [])) > 1
        ]
        if duplicate_crosswalk_pairs:
            crosswalk_conflict = (
                "Multiple crosswalk records target crop labels: "
                + ", ".join(sorted(duplicate_crosswalk_pairs))
            )
            state_audit["conflicting_records"].append(crosswalk_conflict)
            state_audit["report_status"] = "conflicting_records"
        per_state[state] = state_audit

    return {
        "schema_version": 1,
        "source_csv": str(csv_path),
        "report_input_folder": str(report_dir),
        "lgd_download_url": LGD_DOWNLOAD_URL,
        "accepted_report_formats": [".xls", ".xlsx"],
        "state_identity_policy": (
            "The workbook title must identify exactly one 'All Districts of "
            "<state> (State Code:<code>) State' report; its state name must exactly "
            "match a crop CSV state label after whitespace/case normalization. "
            "Filename, district name, and coordinates do not establish report identity."
        ),
        "snapshot_provenance_policy": (
            "source_downloaded_at records the local file modification date, not an "
            "independent download receipt. Each district row retains workbook, "
            "worksheet, row, source URL, LGD state/district codes, and version."
        ),
        "counts": {
            "crop_state_count": len(states),
            "processed_state_count": sum(
                row["report_status"] == "processed"
                for row in per_state.values()
            ),
            "missing_report_count": len(missing_reports),
            "conflicting_state_count": sum(
                row["report_status"] == "conflicting_records"
                for row in per_state.values()
            ),
            "rejected_workbook_count": len(rejected_reports),
            "new_weather_locations_activated": 0,
        },
        "missing_reports": missing_reports,
        "identity_conflicts": identity_conflicts,
        "rejected_reports": rejected_reports,
        "imported_snapshots": imported_snapshots,
        "per_state": per_state,
    }


def crop_pairs(path: Path = CSV_PATH) -> tuple[list[tuple[str, str]], dict[str, set[str]], list[str]]:
    rows = read_csv(path, {"state_names", "district_names"})
    pairs: set[tuple[str, str]] = set()
    districts_by_state: dict[str, set[str]] = defaultdict(set)
    errors: list[str] = []
    for line_number, row in enumerate(rows, start=2):
        state, district = row["state_names"], row["district_names"]
        if not state or not district:
            errors.append(f"CSV row {line_number} has a missing state or district.")
            continue
        pairs.add((state, district))
        districts_by_state[district].add(state)
    return sorted(pairs), districts_by_state, errors


def audit_current_registry(
    csv_path: Path = CSV_PATH,
    registry_path: Path = CURRENT_REGISTRY_PATH,
) -> dict[str, Any]:
    """Create a pair-by-pair audit without promoting any new coordinates."""
    pairs, districts_by_state, source_errors = crop_pairs(csv_path)
    raw_rows = read_csv(csv_path, {"state_names", "district_names"})
    spelling_groups: dict[tuple[str, str], set[tuple[str, str]]] = defaultdict(set)
    for row in raw_rows:
        if row["state_names"] and row["district_names"]:
            normalized_pair = (
                name_key(row["state_names"]),
                name_key(row["district_names"]),
            )
            spelling_groups[normalized_pair].add(
                (row["state_names"], row["district_names"])
            )
    spelling_conflicts = [
        sorted(values)
        for values in spelling_groups.values()
        if len(values) > 1
    ]
    payload = json.loads(registry_path.read_text(encoding="utf-8"))
    registry_rows = payload.get("entries", [])
    existing: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in registry_rows:
        existing[(row.get("state"), row.get("district"))].append(row)

    records: list[dict[str, str]] = []
    for state, district in pairs:
        entries = existing.get((state, district), [])
        entry = entries[0] if len(entries) == 1 else None
        ambiguous = len(districts_by_state[district]) > 1
        if len(entries) > 1:
            status = "suspicious"
            reason = "Registry contains duplicate entries for the exact CSV pair."
        elif entry is None:
            status = "suspicious"
            reason = "CSV pair is missing from the generated registry."
        elif entry.get("coordinate_status") == "verified_reference_point":
            required_fields = (
                "coordinate_source",
                "coordinate_source_url",
                "coordinate_source_record",
                "verification_source_url",
                "verification_note",
            )
            missing_fields = [field for field in required_fields if not entry.get(field)]
            try:
                latitude = _coordinate_value(entry.get("latitude"), "latitude", str((state, district)))
                longitude = _coordinate_value(entry.get("longitude"), "longitude", str((state, district)))
            except ValueError as exc:
                status, reason = "suspicious", str(exc)
            else:
                if missing_fields:
                    status, reason = "suspicious", (
                        "Verified registry record is missing evidence: "
                        + ", ".join(missing_fields)
                    )
                else:
                    status = "verified"
                    reason = (
                        "Existing reviewed pilot reference point; this does not establish "
                        "an LGD district identifier or district-wide coordinate."
                    )
        elif ambiguous:
            status = "ambiguous"
            reason = (
                "The district label occurs in multiple CSV states and has no "
                "reviewed state-qualified administrative crosswalk."
            )
        elif entry.get("coordinate_status") == "unresolved":
            status = "unmatched"
            reason = "No reviewed administrative crosswalk and coordinate evidence."
        else:
            status = "suspicious"
            reason = "Registry has an unrecognized coordinate status."
        records.append(
            {
                "state": state,
                "district": district,
                "status": status,
                "reason": reason,
                "administrative_identifier": (
                    entry.get("administrative_identifier") if entry else None
                ),
                "administrative_state_identifier": (
                    entry.get("administrative_state_identifier") if entry else None
                ),
                "identity_source": entry.get("identity_source") if entry else None,
                "identity_source_url": entry.get("identity_source_url") if entry else None,
                "coordinate_source": entry.get("coordinate_source") if entry else None,
                "coordinate_source_url": (
                    entry.get("coordinate_source_url") if entry else None
                ),
                "coordinate_source_record": (
                    entry.get("coordinate_source_record") if entry else None
                ),
                "latitude": entry.get("latitude") if entry else None,
                "longitude": entry.get("longitude") if entry else None,
                "reference_place": entry.get("reference_place") if entry else None,
                "reference_type": entry.get("reference_type") if entry else None,
                "verification_source_url": (
                    entry.get("verification_source_url") if entry else None
                ),
                "verification_note": entry.get("verification_note") if entry else None,
                "verified_at": entry.get("verified_at") if entry else None,
                "verified_by": entry.get("verified_by") if entry else None,
                "match_method": entry.get("match_method") if entry else None,
                "verification_status": (
                    entry.get("verification_status", entry.get("coordinate_status"))
                    if entry
                    else None
                ),
                "unresolved_reason": (
                    entry.get("unresolved_reason", reason) if entry else reason
                ),
            }
        )

    counts = Counter(record["status"] for record in records)
    return {
        "schema_version": 1,
        "source_csv": str(csv_path.relative_to(BACKEND_ROOT))
        if csv_path.is_relative_to(BACKEND_ROOT)
        else str(csv_path),
        "registry": str(registry_path.relative_to(BACKEND_ROOT))
        if registry_path.is_relative_to(BACKEND_ROOT)
        else str(registry_path),
        "matching_policy": (
            "Exact CSV state/district pair. New matches require an explicit "
            "state-and-LGD-ID crosswalk and reviewed coordinate evidence; "
            "district-name-only and fuzzy matching never promote coordinates."
        ),
        "counts": {
            "supported_pairs": len(pairs),
            "state_count": len({state for state, _ in pairs}),
            "district_label_count": len({district for _, district in pairs}),
            "verified": counts["verified"],
            "ambiguous": counts["ambiguous"],
            "unmatched": counts["unmatched"],
            "suspicious": counts["suspicious"],
            "source_missing_values": len(source_errors),
            "normalized_spelling_conflicts": len(spelling_conflicts),
        },
        "source_quality": {
            "missing_or_blank_labels": source_errors,
            "normalized_spelling_conflicts": spelling_conflicts,
            "district_labels_in_multiple_states": {
                name: sorted(states)
                for name, states in sorted(districts_by_state.items())
                if len(states) > 1
            },
        },
        "source_issues": source_errors,
        "examples": {
            category: [row for row in records if row["status"] == category][:10]
            for category in ("verified", "ambiguous", "unmatched", "suspicious")
        },
        "records": records,
    }


def write_input_templates(output_dir: Path = INPUT_TEMPLATE_DIR) -> dict[str, Path]:
    """Write source-ready templates without inventing administrative matches."""
    report = audit_current_registry()
    output_dir.mkdir(parents=True, exist_ok=True)
    crosswalk_file = output_dir / "crop_to_lgd_crosswalk.csv"
    existing_crosswalk: dict[tuple[str, str], dict[str, str]] = {}
    if crosswalk_file.is_file():
        for existing_row in read_csv(crosswalk_file, CROSSWALK_FIELDS):
            identity = (
                existing_row["crop_state"],
                existing_row["crop_district"],
            )
            if identity in existing_crosswalk:
                raise ValueError(f"Duplicate input crosswalk pair: {identity}")
            existing_crosswalk[identity] = existing_row

    state_file = output_dir / "required_lgd_states.csv"
    states = sorted(
        {
            row["state"]
            for row in report["records"]
            if row["status"] != "verified"
        }
    )
    with state_file.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(
            target,
            fieldnames=[
                "crop_state",
                "state_lgd_code",
                "official_lgd_state_name",
                "source_url",
                "source_record_id",
                "source_downloaded_at",
                "review_note",
            ],
        )
        writer.writeheader()
        for state in states:
            writer.writerow(
                {
                    "crop_state": state,
                    "review_note": (
                        "Fill only from the official LGD state/district directory "
                        "snapshot; preserve the crop label unchanged."
                    ),
                }
            )

    with crosswalk_file.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(
            target,
            fieldnames=sorted(CROSSWALK_FIELDS),
            extrasaction="ignore",
        )
        writer.writeheader()
        for row in report["records"]:
            if row["status"] == "verified":
                continue
            existing_row = existing_crosswalk.get(
                (row["state"], row["district"])
            )
            if existing_row is not None:
                writer.writerow(existing_row)
                continue
            ambiguous = row["status"] == "ambiguous"
            writer.writerow(
                {
                    "crop_state": row["state"],
                    "crop_district": row["district"],
                    "state_lgd_code": "",
                    "district_lgd_code": "",
                    "official_lgd_district_name": "",
                    "district_version": "",
                    "source_record_id": "",
                    "match_method": "",
                    "review_status": "ambiguous" if ambiguous else "pending",
                    "evidence_url": "",
                    "review_note": (
                        "Repeated district label: resolve by state-qualified LGD "
                        "identity before approval; do not choose by name alone."
                        if ambiguous
                        else "Pending state-qualified match to an official LGD record."
                    ),
                    "reviewed_at": "",
                    "reviewed_by": "",
                }
            )

    lgd_file = output_dir / "lgd_district_snapshot_template.csv"
    with lgd_file.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=sorted(LGD_FIELDS))
        writer.writeheader()

    coordinate_file = output_dir / "coordinate_evidence_template.csv"
    with coordinate_file.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=sorted(COORDINATE_FIELDS))
        writer.writeheader()

    return {
        "required_states": state_file,
        "crosswalk": crosswalk_file,
        "lgd_snapshot": lgd_file,
        "coordinate_evidence": coordinate_file,
    }


def _required_review(value: str, field: str, context: str) -> None:
    if not value:
        raise ValueError(f"{context}: required reviewed field {field!r} is empty.")


def _coordinate_value(value: str, field: str, context: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{context}: {field} is not numeric.") from exc
    if not math.isfinite(result):
        raise ValueError(f"{context}: {field} is not finite.")
    limit = 90 if field == "latitude" else 180
    if not -limit <= result <= limit:
        raise ValueError(f"{context}: {field} is outside its geographic range.")
    india_bounds = {
        "latitude": (6.0, 38.0),
        "longitude": (67.0, 98.0),
    }
    lower, upper = india_bounds[field]
    if not lower <= result <= upper:
        raise ValueError(f"{context}: {field} falls outside the India sanity bounds.")
    return result


def build_reviewed_registry(
    lgd_path: Path,
    crosswalk_path: Path,
    coordinate_path: Path,
    csv_path: Path = CSV_PATH,
    current_registry_path: Path = CURRENT_REGISTRY_PATH,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Build a candidate registry only from reviewed, state-qualified evidence."""
    pairs, districts_by_state, source_errors = crop_pairs(csv_path)
    if source_errors:
        raise ValueError("Crop CSV contains missing location labels: " + "; ".join(source_errors[:5]))

    lgd_rows = read_csv(lgd_path, LGD_FIELDS)
    lgd_index: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    lgd_district_code_index: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in lgd_rows:
        _required_review(row["state_lgd_code"], "state_lgd_code", "LGD snapshot")
        _required_review(row["district_lgd_code"], "district_lgd_code", "LGD snapshot")
        _required_review(row["district_version"], "district_version", "LGD snapshot")
        _required_review(row["source_record_id"], "source_record_id", "LGD snapshot")
        _required_review(row["state_name"], "state_name", "LGD snapshot")
        _required_review(row["district_name"], "district_name", "LGD snapshot")
        _required_review(row["source_url"], "source_url", "LGD snapshot")
        _required_review(
            row["source_report_generated_at"],
            "source_report_generated_at",
            "LGD snapshot",
        )
        _required_review(row["source_file"], "source_file", "LGD snapshot")
        _required_review(row["source_worksheet"], "source_worksheet", "LGD snapshot")
        _required_review(
            row["source_worksheet_row"], "source_worksheet_row", "LGD snapshot"
        )
        _required_review(
            row["source_downloaded_at"], "source_downloaded_at", "LGD snapshot"
        )
        try:
            date.fromisoformat(row["source_downloaded_at"][:10])
        except ValueError as exc:
            raise ValueError(
                "LGD snapshot source_downloaded_at must begin with an ISO date."
            ) from exc
        lgd_index[(row["state_lgd_code"], row["district_lgd_code"])].append(row)
        lgd_district_code_index[row["district_lgd_code"]].append(row)

    crosswalk_rows = read_csv(crosswalk_path, CROSSWALK_FIELDS)
    crosswalk_index: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in crosswalk_rows:
        crosswalk_index[(row["crop_state"], row["crop_district"])].append(row)

    coordinate_rows = read_csv(coordinate_path, COORDINATE_FIELDS)
    coordinates_index: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in coordinate_rows:
        coordinates_index[(row["state_lgd_code"], row["district_lgd_code"])].append(row)

    current_payload = json.loads(
        current_registry_path.read_text(encoding="utf-8")
    )
    current_entries: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in current_payload.get("entries", []):
        current_entries[(row.get("state"), row.get("district"))].append(row)

    output_entries: list[dict[str, Any]] = []
    audit_records: list[dict[str, str]] = []
    for state, district in pairs:
        key = (state, district)
        reason: str | None = None
        identity_row: dict[str, str] | None = None
        coordinate_row: dict[str, str] | None = None
        preserved_pilot: dict[str, Any] | None = None
        status = "unmatched"
        mappings = crosswalk_index.get(key, [])

        if len(mappings) > 1:
            status, reason = "suspicious", "Multiple crosswalk rows target the same CSV pair."
        elif not mappings:
            existing_pilots = current_entries.get(key, [])
            if (
                len(existing_pilots) == 1
                and existing_pilots[0].get("coordinate_status")
                == "verified_reference_point"
            ):
                candidate = existing_pilots[0]
                try:
                    _coordinate_value(candidate.get("latitude"), "latitude", str(key))
                    _coordinate_value(candidate.get("longitude"), "longitude", str(key))
                    for field in (
                        "coordinate_source",
                        "coordinate_source_url",
                        "coordinate_source_record",
                        "verification_source_url",
                        "verification_note",
                    ):
                        _required_review(candidate.get(field, ""), field, str(key))
                except ValueError as exc:
                    status, reason = "suspicious", str(exc)
                else:
                    preserved_pilot = candidate
                    status = "verified"
                    reason = (
                        "Preserved existing reviewed pilot reference point; "
                        "it is not an LGD district-wide coordinate."
                    )
            elif len(districts_by_state[district]) > 1:
                status, reason = "ambiguous", "Repeated district label has no state-qualified LGD mapping."
            else:
                reason = "No reviewed crosswalk maps this exact CSV pair to an LGD ID."
        else:
            mapping = mappings[0]
            review_status = mapping["review_status"].casefold()
            if review_status == "approved" and (
                not mapping["reviewed_at"] or not mapping["reviewed_by"]
            ):
                status, reason = "suspicious", "Approved crosswalk lacks reviewer or review date."
            elif review_status == "approved":
                try:
                    date.fromisoformat(mapping["reviewed_at"][:10])
                except ValueError:
                    status, reason = "suspicious", "Approved crosswalk review date must begin with an ISO date."
                else:
                    source_key = (mapping["state_lgd_code"], mapping["district_lgd_code"])
                    source_matches = lgd_index.get(source_key, [])
                    district_id_matches = lgd_district_code_index.get(
                        mapping["district_lgd_code"], []
                    )
                    method = mapping["match_method"].casefold()
                    if len(source_matches) != 1 or len(district_id_matches) != 1:
                        status, reason = "suspicious", "Crosswalk LGD IDs are absent or duplicated in the snapshot."
                    elif method not in {"exact_name", "reviewed_alias", "historical_rename"}:
                        status, reason = "suspicious", "Crosswalk has an unsupported match method."
                    else:
                        identity_row = source_matches[0]
                        record_details_match = (
                            name_key(mapping["official_lgd_district_name"])
                            == name_key(identity_row["district_name"])
                            and mapping["district_version"]
                            == identity_row["district_version"]
                            and mapping["source_record_id"]
                            == identity_row["source_record_id"]
                        )
                        exact_names = (
                            name_key(state) == name_key(identity_row["state_name"])
                            and name_key(district) == name_key(identity_row["district_name"])
                        )
                        if not record_details_match:
                            status, reason = (
                                "suspicious",
                                "Crosswalk official name, district version, or source record "
                                "does not agree with the LGD snapshot.",
                            )
                        elif method == "exact_name" and not exact_names:
                            status, reason = "suspicious", "Exact-name method conflicts with LGD names."
                        elif method != "exact_name" and not (
                            mapping["evidence_url"] and mapping["review_note"]
                        ):
                            status, reason = "suspicious", "Alias or historical-name mapping lacks evidence."
                        else:
                            identity_row = source_matches[0]
                            candidate_coordinates = coordinates_index.get(source_key, [])
                            if len(candidate_coordinates) > 1:
                                status, reason = "suspicious", "Multiple coordinate records exist for the same LGD ID."
                            elif not candidate_coordinates:
                                reason = "LGD identity is mapped, but coordinate evidence is missing."
                            else:
                                candidate = candidate_coordinates[0]
                                required_evidence = (
                                    "reference_place",
                                    "reference_type",
                                    "coordinate_source",
                                    "coordinate_source_url",
                                    "coordinate_source_record",
                                    "verification_source_url",
                                    "verification_note",
                                    "verified_at",
                                    "verified_by",
                                )
                                missing_evidence = [
                                    field for field in required_evidence if not candidate[field]
                                ]
                                if candidate["verification_status"] != "verified":
                                    reason = "Coordinate evidence is not marked verified."
                                elif candidate["reference_type"] != "district_headquarters":
                                    status, reason = "suspicious", (
                                        "Only a source-verified district headquarters "
                                        "coordinate can be promoted; centroids and other "
                                        "reference points are not interchangeable."
                                    )
                                elif missing_evidence:
                                    status, reason = "suspicious", (
                                        "Coordinate evidence is missing: "
                                        + ", ".join(missing_evidence)
                                    )
                                else:
                                    try:
                                        date.fromisoformat(candidate["verified_at"][:10])
                                    except ValueError:
                                        status, reason = "suspicious", (
                                            "Coordinate evidence verified_at must begin with an ISO date."
                                        )
                                    else:
                                        try:
                                            latitude = _coordinate_value(candidate["latitude"], "latitude", str(key))
                                            longitude = _coordinate_value(candidate["longitude"], "longitude", str(key))
                                        except ValueError as exc:
                                            status, reason = "suspicious", str(exc)
                                        else:
                                            coordinate_row = candidate
                                            status = "verified"
                                            reason = None
            elif review_status == "ambiguous":
                status, reason = "ambiguous", mapping["review_note"] or "Mapping is marked ambiguous."
            elif review_status == "suspicious":
                status, reason = "suspicious", mapping["review_note"] or "Mapping is marked suspicious."
            elif review_status != "approved":
                reason = mapping["review_note"] or "Mapping has not been approved."

        if status == "verified" and identity_row and coordinate_row:
            entry: dict[str, Any] = {
                "state": state,
                "district": district,
                "administrative_identifier": identity_row["district_lgd_code"],
                "administrative_state_identifier": identity_row["state_lgd_code"],
                "official_lgd_district_name": identity_row["district_name"],
                "district_version": identity_row["district_version"],
                "identity_source": "Local Government Directory (LGD)",
                "identity_source_url": identity_row["source_url"],
                "identity_source_record": identity_row["source_record_id"],
                "match_method": mappings[0]["match_method"],
                "coordinate_status": "verified_reference_point",
                "latitude": latitude,
                "longitude": longitude,
                "reference_place": coordinate_row["reference_place"],
                "reference_type": coordinate_row["reference_type"],
                "coordinate_source": coordinate_row["coordinate_source"],
                "coordinate_source_url": coordinate_row["coordinate_source_url"],
                "coordinate_source_record": coordinate_row["coordinate_source_record"],
                "verification_source_url": coordinate_row["verification_source_url"],
                "verification_note": coordinate_row["verification_note"],
                "verification_status": "verified",
                "verified_at": coordinate_row["verified_at"],
                "verified_by": coordinate_row["verified_by"],
                "unresolved_reason": None,
            }
        elif status == "verified" and preserved_pilot:
            entry = dict(preserved_pilot)
            entry["unresolved_reason"] = None
        else:
            entry = {
                "state": state,
                "district": district,
                "administrative_state_identifier": (
                    identity_row["state_lgd_code"] if identity_row else None
                ),
                "administrative_identifier": (
                    identity_row["district_lgd_code"] if identity_row else None
                ),
                "official_lgd_district_name": (
                    identity_row["district_name"] if identity_row else None
                ),
                "district_version": (
                    identity_row["district_version"] if identity_row else None
                ),
                "identity_source": (
                    "Local Government Directory (LGD)" if identity_row else None
                ),
                "identity_source_url": (
                    identity_row["source_url"] if identity_row else None
                ),
                "identity_source_record": (
                    identity_row["source_record_id"] if identity_row else None
                ),
                "match_method": mappings[0]["match_method"] if identity_row else None,
                "coordinate_status": "unresolved",
                "latitude": None,
                "longitude": None,
                "verification_status": status,
                "unresolved_reason": reason,
            }
        output_entries.append(entry)
        audit_records.append(
            {
                "state": state,
                "district": district,
                "status": status,
                "reason": reason or "",
                "administrative_identifier": entry.get("administrative_identifier"),
                "official_lgd_district_name": entry.get(
                    "official_lgd_district_name"
                ),
                "district_version": entry.get("district_version"),
                "administrative_state_identifier": entry.get(
                    "administrative_state_identifier"
                ),
                "identity_source": entry.get("identity_source"),
                "identity_source_url": entry.get("identity_source_url"),
                "identity_source_record": entry.get("identity_source_record"),
                "coordinate_source": entry.get("coordinate_source"),
                "coordinate_source_url": entry.get("coordinate_source_url"),
                "coordinate_source_record": entry.get("coordinate_source_record"),
                "latitude": entry.get("latitude"),
                "longitude": entry.get("longitude"),
                "reference_place": entry.get("reference_place"),
                "reference_type": entry.get("reference_type"),
                "verification_source_url": entry.get("verification_source_url"),
                "verification_note": entry.get("verification_note"),
                "verified_at": entry.get("verified_at"),
                "verified_by": entry.get("verified_by"),
                "match_method": entry.get("match_method"),
                "verification_status": entry.get("verification_status"),
                "unresolved_reason": entry.get("unresolved_reason"),
            }
        )

    counts = Counter(row["status"] for row in audit_records)
    registry = {
        "schema_version": 2,
        "source_csv": csv_path.name,
        "identity_source": "Local Government Directory (LGD) snapshot",
        "coordinate_source_note": (
            "Coordinates are included only when coordinate evidence is explicitly "
            "reviewed and tied to the mapped LGD state/district identifiers."
        ),
        "supported_pair_count": len(output_entries),
        "verified_pair_count": counts["verified"],
        "unresolved_pair_count": len(output_entries) - counts["verified"],
        "entries": output_entries,
    }
    audit = {
        "schema_version": 1,
        "matching_policy": (
            "Explicit state-and-district LGD identifier crosswalk; no fuzzy "
            "or district-name-only matching."
        ),
        "counts": {
            "supported_pairs": len(audit_records),
            "verified": counts["verified"],
            "ambiguous": counts["ambiguous"],
            "unmatched": counts["unmatched"],
            "suspicious": counts["suspicious"],
        },
        "examples": {
            category: [
                row for row in audit_records if row["status"] == category
            ][:10]
            for category in ("verified", "ambiguous", "unmatched", "suspicious")
        },
        "records": audit_records,
    }
    return registry, audit


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit-current", action="store_true")
    parser.add_argument("--audit-lgd-reports", action="store_true")
    parser.add_argument("--prepare-inputs", action="store_true")
    parser.add_argument("--template-dir", type=Path, default=INPUT_TEMPLATE_DIR)
    parser.add_argument("--report-dir", type=Path, default=LGD_REPORT_INPUT_DIR)
    parser.add_argument("--batch-audit-output", type=Path, default=BATCH_AUDIT_PATH)
    parser.add_argument("--lgd-snapshot", type=Path)
    parser.add_argument("--crosswalk", type=Path)
    parser.add_argument("--coordinate-evidence", type=Path)
    parser.add_argument("--registry-output", type=Path)
    parser.add_argument("--audit-output", type=Path, default=AUDIT_PATH)
    args = parser.parse_args()

    if args.audit_current:
        report = audit_current_registry()
        args.audit_output.parent.mkdir(parents=True, exist_ok=True)
        args.audit_output.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(json.dumps(report["counts"], sort_keys=True))
        return
    if args.audit_lgd_reports:
        report = audit_lgd_reports(
            report_dir=args.report_dir,
            output_dir=args.template_dir,
        )
        args.batch_audit_output.parent.mkdir(parents=True, exist_ok=True)
        args.batch_audit_output.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(json.dumps(report["counts"], sort_keys=True))
        return
    if args.prepare_inputs:
        paths = write_input_templates(args.template_dir)
        print(
            json.dumps(
                {name: str(path) for name, path in paths.items()},
                sort_keys=True,
            )
        )
        return

    required = {
        "--lgd-snapshot": args.lgd_snapshot,
        "--crosswalk": args.crosswalk,
        "--coordinate-evidence": args.coordinate_evidence,
        "--registry-output": args.registry_output,
    }
    missing = [name for name, value in required.items() if value is None]
    if missing:
        parser.error("reviewed-registry mode requires " + ", ".join(missing))

    registry, audit = build_reviewed_registry(
        args.lgd_snapshot,
        args.crosswalk,
        args.coordinate_evidence,
    )
    args.registry_output.parent.mkdir(parents=True, exist_ok=True)
    args.audit_output.parent.mkdir(parents=True, exist_ok=True)
    args.registry_output.write_text(
        json.dumps(registry, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    args.audit_output.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(audit["counts"], sort_keys=True))


if __name__ == "__main__":
    main()
