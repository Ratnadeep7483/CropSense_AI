#!/usr/bin/env python3
"""Compare a merged LGD workbook with crop labels and reviewed location evidence."""

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

BACKEND_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = BACKEND_ROOT / "data"
RESOLUTION_DIR = DATA_DIR / "weather_location_resolution"
WORKBOOK_PATH = DATA_DIR / "All_Districts_Merged.xlsx"
CROP_PATH = DATA_DIR / "Copy of crop_recommendation_45_with_year_area.csv"
CROSSWALK_PATH = RESOLUTION_DIR / "crop_to_lgd_crosswalk.csv"
KARNATAKA_SNAPSHOT_PATH = RESOLUTION_DIR / "lgd_district_snapshot_karnataka.csv"
KARNATAKA_AUDIT_PATH = RESOLUTION_DIR / "karnataka_lgd_crosswalk_audit.json"
OUTPUT_PATH = RESOLUTION_DIR / "merged_lgd_workbook_audit.json"

EXPECTED_HEADERS = [
    "S. No.",
    "State Name",
    "State Code",
    "District Code",
    "District Version",
    "District Name\n(In English)",
    "District Name\n(In Local)",
    "Census 2001 Code",
    "Census 2011 Code",
]


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def name_key(value: str) -> str:
    return re.sub(r"\s+", " ", clean(value)).casefold()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def load_workbook_records(
    workbook_path: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Districts" not in workbook.sheetnames:
            raise ValueError("Merged LGD workbook is missing the Districts sheet.")
        worksheet = workbook["Districts"]
        rows = worksheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows, ())]
        if headers != EXPECTED_HEADERS:
            raise ValueError(f"Unexpected merged workbook headers: {headers!r}")

        records: list[dict[str, Any]] = []
        for worksheet_row, values in enumerate(rows, start=2):
            if not any(value is not None for value in values):
                continue
            record = {
                "worksheet_row": worksheet_row,
                "serial_number": clean(values[0]),
                "state_name": clean(values[1]),
                "state_code": clean(values[2]),
                "district_code": clean(values[3]),
                "district_version": clean(values[4]),
                "district_name_english": clean(values[5]),
                "district_name_local": clean(values[6]),
                "census_2001_code": clean(values[7]),
                "census_2011_code": clean(values[8]),
                "source_record_id": (
                    f"{workbook_path.name}:Districts!"
                    f"D{worksheet_row}:F{worksheet_row}"
                ),
            }
            records.append(record)

        properties = workbook.properties
        core_properties = {
            "creator": properties.creator,
            "last_modified_by": properties.lastModifiedBy,
            "created": (
                properties.created.isoformat() if properties.created else None
            ),
            "modified": (
                properties.modified.isoformat() if properties.modified else None
            ),
            "title": properties.title,
            "subject": properties.subject,
            "description": properties.description,
            "identifier": properties.identifier,
        }
        return records, {
            "sheet_names": list(workbook.sheetnames),
            "sheet_dimensions": {
                sheet.title: {
                    "rows_including_header": sheet.max_row,
                    "columns": sheet.max_column,
                    "state": sheet.sheet_state,
                }
                for sheet in workbook.worksheets
            },
            "headers": headers,
            "core_properties": core_properties,
        }
    finally:
        workbook.close()


def index_crosswalk(
    rows: list[dict[str, str]],
) -> dict[tuple[str, str], list[dict[str, str]]]:
    result: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        result[(row["crop_state"], row["crop_district"])].append(row)
    return result


def main() -> None:
    records, workbook_metadata = load_workbook_records(WORKBOOK_PATH)
    if len(records) != 781:
        raise ValueError(f"Expected 781 district rows, found {len(records)}.")

    crop_rows = read_csv(CROP_PATH)
    crosswalk_rows = read_csv(CROSSWALK_PATH)
    karnataka_rows = read_csv(KARNATAKA_SNAPSHOT_PATH)
    prior_karnataka_audit = json.loads(
        KARNATAKA_AUDIT_PATH.read_text(encoding="utf-8")
    )
    crop_pairs = sorted(
        {
            (clean(row["state_names"]), clean(row["district_names"]))
            for row in crop_rows
            if clean(row["state_names"]) and clean(row["district_names"])
        }
    )
    crop_states = sorted({state for state, _ in crop_pairs})
    crosswalk_by_pair = index_crosswalk(crosswalk_rows)

    records_by_state: dict[str, list[dict[str, Any]]] = defaultdict(list)
    records_by_id: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    records_by_district_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    records_by_pair: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        records_by_state[name_key(record["state_name"])].append(record)
        records_by_id[(record["state_code"], record["district_code"])].append(
            record
        )
        records_by_district_code[record["district_code"]].append(record)
        records_by_pair[
            (name_key(record["state_name"]), name_key(record["district_name_english"]))
        ].append(record)

    per_state: dict[str, dict[str, Any]] = {}
    for crop_state in crop_states:
        state_records = records_by_state.get(name_key(crop_state), [])
        state_names = sorted({row["state_name"] for row in state_records})
        state_codes = sorted({row["state_code"] for row in state_records})
        official_names: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in state_records:
            official_names[name_key(row["district_name_english"])].append(row)

        crop_districts = sorted(
            district for state, district in crop_pairs if state == crop_state
        )
        direct_matches: list[dict[str, Any]] = []
        reviewed_aliases: list[dict[str, Any]] = []
        unresolved: list[str] = []
        conflicts: list[dict[str, str]] = []
        for crop_district in crop_districts:
            candidates = official_names.get(name_key(crop_district), [])
            mappings = crosswalk_by_pair.get((crop_state, crop_district), [])
            if len(mappings) > 1:
                conflicts.append(
                    {
                        "crop_district": crop_district,
                        "reason": "Duplicate crosswalk rows target this pair.",
                    }
                )

            if len(candidates) == 1:
                candidate = candidates[0]
                direct_matches.append(
                    {
                        "crop_district": crop_district,
                        "official_district_name": candidate[
                            "district_name_english"
                        ],
                        "state_code": candidate["state_code"],
                        "district_code": candidate["district_code"],
                        "district_version": candidate["district_version"],
                        "source_record_id": candidate["source_record_id"],
                        "worksheet_row": candidate["worksheet_row"],
                    }
                )
                if len(mappings) == 1:
                    mapping = mappings[0]
                    if mapping["review_status"].casefold() == "approved":
                        agrees = (
                            mapping["match_method"].casefold() == "exact_name"
                            and mapping["state_lgd_code"] == candidate["state_code"]
                            and mapping["district_lgd_code"]
                            == candidate["district_code"]
                            and name_key(mapping["official_lgd_district_name"])
                            == name_key(candidate["district_name_english"])
                            and mapping["district_version"]
                            == candidate["district_version"]
                        )
                        if not agrees:
                            conflicts.append(
                                {
                                    "crop_district": crop_district,
                                    "reason": (
                                        "Approved exact-name crosswalk identity "
                                        "does not agree with merged workbook."
                                    ),
                                }
                            )
                continue

            if len(candidates) > 1:
                unresolved.append(crop_district)
                conflicts.append(
                    {
                        "crop_district": crop_district,
                        "reason": (
                            "Multiple workbook rows share this state-qualified "
                            "English district name."
                        ),
                    }
                )
                continue

            if len(mappings) == 1:
                mapping = mappings[0]
                method = mapping["match_method"].casefold()
                targets = records_by_id.get(
                    (mapping["state_lgd_code"], mapping["district_lgd_code"]), []
                )
                if (
                    mapping["review_status"].casefold() == "approved"
                    and method in {"reviewed_alias", "historical_rename"}
                    and len(targets) == 1
                    and name_key(mapping["official_lgd_district_name"])
                    == name_key(targets[0]["district_name_english"])
                    and mapping["district_version"]
                    == targets[0]["district_version"]
                    and mapping["evidence_url"]
                    and mapping["review_note"]
                    and targets[0]["state_code"] in state_codes
                ):
                    reviewed_aliases.append(
                        {
                            "crop_district": crop_district,
                            "official_district_name": targets[0][
                                "district_name_english"
                            ],
                            "state_code": targets[0]["state_code"],
                            "district_code": targets[0]["district_code"],
                            "district_version": targets[0]["district_version"],
                            "evidence_url": mapping["evidence_url"],
                            "reviewed_at": mapping["reviewed_at"],
                            "reviewed_by": mapping["reviewed_by"],
                        }
                    )
                    continue
            unresolved.append(crop_district)

        per_state[crop_state] = {
            "crop_state_label": crop_state,
            "workbook_state_name": state_names[0] if len(state_names) == 1 else None,
            "state_label_match": (
                "case_whitespace_normalized" if state_records else "missing"
            ),
            "state_code": state_codes[0] if len(state_codes) == 1 else None,
            "state_identity_conflicts": (
                []
                if len(state_names) <= 1 and len(state_codes) <= 1
                else {"workbook_state_names": state_names, "state_codes": state_codes}
            ),
            "crop_pair_count": len(crop_districts),
            "lgd_record_count": len(state_records),
            "direct_match_count": len(direct_matches),
            "direct_matches": direct_matches,
            "reviewed_alias_count": len(reviewed_aliases),
            "reviewed_aliases": reviewed_aliases,
            "unresolved_count": len(unresolved),
            "unresolved_names": unresolved,
            "conflicting_records": conflicts,
        }

    required_fields = (
        "state_name",
        "state_code",
        "district_code",
        "district_version",
        "district_name_english",
    )
    missing_fields = [
        {
            "worksheet_row": row["worksheet_row"],
            "fields": [field for field in required_fields if not row[field]],
        }
        for row in records
        if any(not row[field] for field in required_fields)
    ]
    invalid_identifiers = [
        {
            "worksheet_row": row["worksheet_row"],
            "field": field,
            "value": row[field],
        }
        for row in records
        for field in ("state_code", "district_code", "district_version")
        if not row[field].isdigit()
    ]
    duplicate_pairs = [
        {
            "state": rows[0]["state_name"],
            "district_name_english": rows[0]["district_name_english"],
            "worksheet_rows": [row["worksheet_row"] for row in rows],
        }
        for rows in records_by_pair.values()
        if len(rows) > 1
    ]
    duplicate_codes = [
        {
            "district_code": code,
            "records": [
                {
                    "state": row["state_name"],
                    "state_code": row["state_code"],
                    "district": row["district_name_english"],
                    "district_version": row["district_version"],
                    "worksheet_row": row["worksheet_row"],
                }
                for row in rows
            ],
        }
        for code, rows in records_by_id.items()
        if len(rows) > 1
    ]
    duplicate_district_codes = [
        {
            "district_code": code,
            "records": [
                {
                    "state": row["state_name"],
                    "state_code": row["state_code"],
                    "district": row["district_name_english"],
                    "district_version": row["district_version"],
                    "worksheet_row": row["worksheet_row"],
                }
                for row in rows
            ],
        }
        for code, rows in records_by_district_code.items()
        if len(rows) > 1
    ]

    state_code_by_name: dict[str, set[str]] = defaultdict(set)
    state_names_by_code: dict[str, set[str]] = defaultdict(set)
    for row in records:
        state_code_by_name[row["state_name"]].add(row["state_code"])
        state_names_by_code[row["state_code"]].add(row["state_name"])
    state_name_code_conflicts = {
        name: sorted(codes)
        for name, codes in state_code_by_name.items()
        if len(codes) != 1
    }
    state_code_name_conflicts = {
        code: sorted(names)
        for code, names in state_names_by_code.items()
        if len(names) != 1
    }

    karnataka_by_id = {
        (row["state_lgd_code"], row["district_lgd_code"]): row
        for row in karnataka_rows
    }
    karnataka_records = [
        row
        for row in records
        if name_key(row["state_name"]) == name_key("Karnataka")
    ]
    karnataka_discrepancies = []
    for row in karnataka_records:
        prior = karnataka_by_id.get((row["state_code"], row["district_code"]))
        if prior is None:
            karnataka_discrepancies.append(
                {
                    "type": "merged_record_missing_from_prior_snapshot",
                    "district_code": row["district_code"],
                    "district_name": row["district_name_english"],
                }
            )
        elif (
            name_key(prior["district_name"])
            != name_key(row["district_name_english"])
            or prior["district_version"] != row["district_version"]
        ):
            karnataka_discrepancies.append(
                {
                    "type": "name_or_version_difference",
                    "district_code": row["district_code"],
                    "prior_name": prior["district_name"],
                    "merged_name": row["district_name_english"],
                    "prior_version": prior["district_version"],
                    "merged_version": row["district_version"],
                }
            )
    for (state_code, district_code), prior in karnataka_by_id.items():
        if not any(
            row["state_code"] == state_code
            and row["district_code"] == district_code
            for row in karnataka_records
        ):
            karnataka_discrepancies.append(
                {
                    "type": "prior_snapshot_record_missing_from_merged",
                    "district_code": district_code,
                    "district_name": prior["district_name"],
                }
            )

    approved_crosswalk_comparison = []
    for pair, mappings in crosswalk_by_pair.items():
        for mapping in mappings:
            if mapping["review_status"].casefold() != "approved":
                continue
            targets = records_by_id.get(
                (mapping["state_lgd_code"], mapping["district_lgd_code"]), []
            )
            if len(targets) != 1:
                status = "ids_missing_or_duplicated"
                target = None
            else:
                target = targets[0]
                status = (
                    "agrees"
                    if (
                        name_key(mapping["official_lgd_district_name"])
                        == name_key(target["district_name_english"])
                        and mapping["district_version"]
                        == target["district_version"]
                        and mapping["state_lgd_code"] == target["state_code"]
                        and name_key(mapping["crop_state"])
                        == name_key(target["state_name"])
                    )
                    else "identity_mismatch"
                )
            approved_crosswalk_comparison.append(
                {
                    "crop_state": pair[0],
                    "crop_district": pair[1],
                    "review_status": mapping["review_status"],
                    "match_method": mapping["match_method"],
                    "result": status,
                    "crosswalk_source_record_id": mapping["source_record_id"],
                    "merged_source_record_id": (
                        target["source_record_id"] if target else None
                    ),
                    "note": (
                        "Different source record IDs are expected for independent "
                        "workbooks; identity comparison uses state/district codes, "
                        "official English name, and version."
                    ),
                }
            )

    missing_crop_states = [
        state
        for state in crop_states
        if per_state[state]["state_label_match"] == "missing"
    ]
    crop_state_keys = {name_key(state) for state in crop_states}
    extra_workbook_states = sorted(
        {
            row["state_name"]
            for row in records
            if name_key(row["state_name"]) not in crop_state_keys
        }
    )
    direct_total = sum(
        state["direct_match_count"] for state in per_state.values()
    )
    aliases_total = sum(
        state["reviewed_alias_count"] for state in per_state.values()
    )
    unresolved_total = sum(
        state["unresolved_count"] for state in per_state.values()
    )
    report = {
        "audit_only": True,
        "created_at_local": datetime.now().astimezone().isoformat(timespec="seconds"),
        "workbook": {
            "path": str(WORKBOOK_PATH),
            "sha256": hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest(),
            "size_bytes": WORKBOOK_PATH.stat().st_size,
            "data_records": len(records),
            **workbook_metadata,
            "source_provenance_assessment": (
                "No publisher/source URL or provenance evidence was found in "
                "the workbook core properties. Creator metadata is not proof of "
                "origin; do not treat the filename or LGD-like content as official "
                "authentication."
            ),
        },
        "comparison_policy": (
            "State and English district names compared case-insensitively with "
            "whitespace normalization. District matching is state-qualified. No "
            "fuzzy matching, alias inference, code inference, or coordinate inference. "
            "Existing approved crosswalk aliases are recognized only when their "
            "state/district LGD IDs, official English name, version, and evidence "
            "agree with a unique workbook record."
        ),
        "inputs": {
            "crop_csv": str(CROP_PATH),
            "crosswalk": str(CROSSWALK_PATH),
            "karnataka_snapshot": str(KARNATAKA_SNAPSHOT_PATH),
            "prior_karnataka_audit": str(KARNATAKA_AUDIT_PATH),
        },
        "counts": {
            "crop_state_labels": len(crop_states),
            "crop_state_district_pairs": len(crop_pairs),
            "workbook_state_count": len(
                {name_key(row["state_name"]) for row in records}
            ),
            "crop_states_matching_workbook": len(crop_states)
            - len(missing_crop_states),
            "crop_states_missing_from_workbook": len(missing_crop_states),
            "direct_district_matches": direct_total,
            "reviewed_alias_matches": aliases_total,
            "unresolved_crop_district_names": unresolved_total,
            "missing_required_identifier_or_name_records": len(missing_fields),
            "invalid_numeric_identifier_values": len(invalid_identifiers),
            "duplicate_state_district_pairs": len(duplicate_pairs),
            "duplicate_state_qualified_district_codes": len(duplicate_codes),
            "duplicate_district_codes_globally": len(duplicate_district_codes),
            "state_names_with_multiple_state_codes": len(
                state_name_code_conflicts
            ),
            "state_codes_with_multiple_state_names": len(
                state_code_name_conflicts
            ),
            "approved_crosswalk_rows_compared": len(
                approved_crosswalk_comparison
            ),
            "approved_crosswalk_identity_mismatches": sum(
                row["result"] != "agrees"
                for row in approved_crosswalk_comparison
            ),
            "karnataka_snapshot_discrepancies": len(karnataka_discrepancies),
            "weather_locations_activated": 0,
        },
        "state_label_matches": {
            state: {
                "workbook_state_name": per_state[state]["workbook_state_name"],
                "state_code": per_state[state]["state_code"],
                "match": per_state[state]["state_label_match"],
            }
            for state in crop_states
        },
        "missing_crop_states": missing_crop_states,
        "workbook_states_not_in_crop_dataset": extra_workbook_states,
        "quality_checks": {
            "missing_required_fields": missing_fields,
            "invalid_numeric_identifiers": invalid_identifiers,
            "duplicate_state_district_pairs": duplicate_pairs,
            "duplicate_state_qualified_district_codes": duplicate_codes,
            "duplicate_district_codes_globally": duplicate_district_codes,
            "state_name_code_conflicts": state_name_code_conflicts,
            "state_code_name_conflicts": state_code_name_conflicts,
        },
        "approved_crosswalk_comparison": approved_crosswalk_comparison,
        "karnataka_comparison": {
            "merged_workbook_records": len(karnataka_records),
            "prior_snapshot_records": len(karnataka_rows),
            "prior_audit_direct_matches": prior_karnataka_audit["counts"][
                "direct_normalized_identity_matches"
            ],
            "merged_workbook_direct_matches": per_state["Karnataka"][
                "direct_match_count"
            ],
            "discrepancies": karnataka_discrepancies,
            "result": (
                "No discrepancy in LGD state code, district code, normalized "
                "English district name, or district version among the 31 records."
                if not karnataka_discrepancies
                else "Differences found; prior snapshot was not modified."
            ),
        },
        "per_state": per_state,
        "safety": {
            "source_files_modified": False,
            "weather_registry_modified": False,
            "models_or_apis_modified": False,
            "weather_locations_activated": 0,
        },
    }
    OUTPUT_PATH.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "output": str(OUTPUT_PATH),
                "counts": report["counts"],
                "per_state": {
                    state: {
                        "state_code": item["state_code"],
                        "crop_pairs": item["crop_pair_count"],
                        "direct": item["direct_match_count"],
                        "reviewed_aliases": item["reviewed_alias_count"],
                        "unresolved": item["unresolved_names"],
                    }
                    for state, item in per_state.items()
                },
                "extra_workbook_states": extra_workbook_states,
                "karnataka_discrepancies": karnataka_discrepancies,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
